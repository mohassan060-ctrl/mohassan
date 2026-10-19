"""Review workflow: persistence, evidence provenance, and deterministic drafting."""
import io
import json
import os
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parent
DATA = BASE / 'data'
DATA.mkdir(exist_ok=True)
DB = Path(os.environ.get('VISASPONSOR_DB_PATH', str(DATA / 'applications.sqlite3')))
STAGES = ['New', 'Researching sponsorship', 'Ready for review', 'Approved to apply', 'Applied', 'Interviewing', 'Rejected', 'Offer', 'Archived']
EVIDENCE_KINDS = ['Role posting', 'Employer policy', 'Official sponsor register', 'Employer confirmation', 'Other']
EVIDENCE_VERDICTS = ['Explicit role sponsorship', 'Explicit no sponsorship', 'Employer licensed / past sponsorship only', 'Unclear / indirect', 'Employer confirmed for this vacancy']

def now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')

def conn():
    DB.parent.mkdir(parents=True, exist_ok=True)
    c=sqlite3.connect(DB)
    c.row_factory=sqlite3.Row
    c.execute('PRAGMA journal_mode=WAL')
    c.executescript('''
    CREATE TABLE IF NOT EXISTS jobs (
      job_id TEXT PRIMARY KEY, company TEXT NOT NULL, title TEXT NOT NULL, location TEXT,
      url TEXT, description TEXT, source TEXT, score INTEGER, screening TEXT,
      status TEXT DEFAULT 'New', notes TEXT DEFAULT '', first_seen TEXT, last_seen TEXT,
      approved_at TEXT, applied_at TEXT, package_version INTEGER DEFAULT 0
    );
    CREATE TABLE IF NOT EXISTS evidence (
      id INTEGER PRIMARY KEY AUTOINCREMENT, job_id TEXT NOT NULL, kind TEXT NOT NULL,
      verdict TEXT NOT NULL, quote TEXT NOT NULL, source_url TEXT NOT NULL,
      checked_at TEXT NOT NULL, reviewer TEXT DEFAULT 'Automated extraction',
      FOREIGN KEY(job_id) REFERENCES jobs(job_id)
    );
    CREATE TABLE IF NOT EXISTS decisions (
      id INTEGER PRIMARY KEY AUTOINCREMENT, job_id TEXT NOT NULL, stage TEXT NOT NULL,
      decision_at TEXT NOT NULL, notes TEXT DEFAULT ''
    );
    ''')
    return c

def upsert_jobs(jobs):
    with conn() as c:
        for j in jobs:
            id_=j['job_id']
            existed=c.execute('SELECT job_id FROM jobs WHERE job_id=?',(id_,)).fetchone()
            c.execute('''INSERT INTO jobs(job_id,company,title,location,url,description,source,score,screening,first_seen,last_seen)
               VALUES (?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(job_id) DO UPDATE SET
               company=excluded.company,title=excluded.title,location=excluded.location,
               url=excluded.url,description=excluded.description,source=excluded.source,
               score=excluded.score,screening=excluded.screening,last_seen=excluded.last_seen''',
               (id_,j['company'],j['title'],j.get('location',''),j.get('url',''),j.get('description',''),
                j.get('source',''),int(j.get('match_score',0)),j.get('sponsorship','Not stated'),now(),now()))
            if not existed:
                classification=j.get('sponsorship','Not stated')
                verdict={'Explicitly offered':'Explicit role sponsorship','Unavailable':'Explicit no sponsorship'}.get(classification,'Unclear / indirect')
                c.execute('INSERT INTO evidence(job_id,kind,verdict,quote,source_url,checked_at) VALUES (?,?,?,?,?,?)',
                          (id_,'Role posting',verdict,j.get('evidence','No qualifying sponsorship statement located.'),j.get('url',''),now()))

def all_jobs():
    with conn() as c:
        return [dict(r) for r in c.execute('SELECT * FROM jobs ORDER BY score DESC, last_seen DESC')]

def job_by_id(jid):
    with conn() as c:
        r=c.execute('SELECT * FROM jobs WHERE job_id=?',(jid,)).fetchone()
        return dict(r) if r else None

def get_evidence(jid):
    with conn() as c:
        return [dict(x) for x in c.execute('SELECT * FROM evidence WHERE job_id=? ORDER BY id DESC',(jid,))]

def add_evidence(jid,kind,verdict,quote,url,reviewer='Candidate review'):
    if not quote.strip() or not re.match(r'^https?://',url):
        raise ValueError('Evidence requires a substantive note or quotation and a source URL.')
    with conn() as c:
        c.execute('INSERT INTO evidence(job_id,kind,verdict,quote,source_url,checked_at,reviewer) VALUES (?,?,?,?,?,?,?)',
                  (jid,kind,verdict,quote.strip(),url.strip(),now(),reviewer))

def sponsorship_assessment(evidence):
    """Negative job-level evidence blocks approval until employer explicitly confirms the specific vacancy."""
    confirmations=[e for e in evidence if e['verdict']=='Employer confirmed for this vacancy']
    negatives=[e for e in evidence if e['verdict']=='Explicit no sponsorship']
    positives=[e for e in evidence if e['verdict']=='Explicit role sponsorship']
    if confirmations:
        return 'Employer confirmed for this vacancy', True
    if negatives:
        return 'Excluded by published vacancy', False
    if positives:
        return 'Role posting states sponsorship; verify eligibility', True
    if any(e['verdict']=='Employer licensed / past sponsorship only' for e in evidence):
        return 'Employer record only — role unverified', False
    return 'No verified sponsorship for this vacancy', False

def set_stage(jid, stage, notes=''):
    if stage not in STAGES: raise ValueError('Unknown stage')
    with conn() as c:
        r=c.execute('SELECT status FROM jobs WHERE job_id=?',(jid,)).fetchone()
        if not r: raise ValueError('Unknown job')
        if stage=='Approved to apply':
            ev=[dict(e) for e in c.execute('SELECT * FROM evidence WHERE job_id=?',(jid,))]
            _,ok=sponsorship_assessment(ev)
            if not ok: raise ValueError('Approval blocked: add affirmative evidence that this specific role sponsors immigration permission.')
        if stage=='Applied' and r['status'] not in ('Approved to apply','Applied'):
            raise ValueError('Application cannot be marked Applied until it is approved.')
        fields=['status=?','notes=?']
        params=[stage,notes]
        if stage=='Approved to apply': fields.append('approved_at=?');params.append(now())
        if stage=='Applied': fields.append('applied_at=?');params.append(now())
        params.append(jid)
        c.execute('UPDATE jobs SET '+','.join(fields)+' WHERE job_id=?',params)
        c.execute('INSERT INTO decisions(job_id,stage,decision_at,notes) VALUES (?,?,?,?)',(jid,stage,now(),notes))

def build_brief(job,profile,evidence):
    assessment,_=sponsorship_assessment(evidence)
    return {'candidate_label':profile.get('candidate_label','International applicant'),
            'job':job, 'sponsorship_assessment':assessment,
            'evidence':[{'kind':e['kind'],'verdict':e['verdict'],'quote':e['quote'],'source':e['source_url'],'checked_at':e['checked_at']} for e in evidence],
            'verified_facts':profile.get('verified_achievements',[]),
            'claims_to_confirm':profile.get('claims_to_confirm',[]),
            'disclaimer':'Drafts are not submitted. All statements and immigration eligibility must be reviewed before use.'}

def relevant_facts(job,profile,limit=9):
    desc=(job['title']+' '+job.get('description','')).lower()
    groups={'communications':['communication','pr ','public relations','media','editorial','stakeholder','messaging','thought leadership','internal'],
            'marketing':['marketing','brand','campaign','demand','lead generation','product'],
            'change':['change','consult','organisation','organizational','transformation','workforce']}
    words=set(re.findall(r'[a-z]{4,}',desc))
    ranked=[]
    for i,fact in enumerate(profile.get('verified_achievements',[])):
        if any(x in fact.lower() for x in ['reported in communications-oriented cv','ba (honours)','ms management','certificat','7+ years']):
            continue
        fw=set(re.findall(r'[a-z]{4,}',fact.lower()))
        score=len(words & fw)
        for terms in groups.values():
            if any(t in desc for t in terms) and any(t in fact.lower() for t in terms): score+=3
        ranked.append((score,-i,fact))
    return [item[2] for item in sorted(ranked,reverse=True)[:limit]]

def setup_doc(doc):
    sec=doc.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72)
    sec.left_margin=sec.right_margin=Inches(.8)
    style=doc.styles['Normal'];style.font.name='Aptos';style.font.size=Pt(10)
    style.paragraph_format.space_after=Pt(5)
    for s in ['Title','Heading 1','Heading 2']:
        doc.styles[s].font.name='Aptos';doc.styles[s].font.color.rgb=RGBColor(26,54,85)

def docx_bytes(doc):
    out=io.BytesIO();doc.save(out);return out.getvalue()

def make_cv_docx(job, profile, contact=None):
    contact=contact or {}
    doc=Document();setup_doc(doc)
    doc.add_heading(contact.get('name','SYED MOHAMMAD HASSAN'),0)
    doc.add_paragraph(' | '.join(x for x in [contact.get('location','Karachi, Pakistan'),contact.get('email',''),contact.get('linkedin','')] if x))
    doc.add_heading('TARGET ROLE',level=1)
    doc.add_paragraph(job['title']+' | '+job['company'])
    doc.add_heading('PROFESSIONAL PROFILE',level=1)
    doc.add_paragraph('Communications and marketing professional with over seven years of experience spanning regional brand communications, organisational change consulting and editorial work. Experience includes Dun & Bradstreet, PwC and The Express Tribune.')
    doc.add_heading('RELEVANT EXPERIENCE HIGHLIGHTS',level=1)
    for fact in relevant_facts(job,profile):doc.add_paragraph(fact,style='List Bullet')
    doc.add_heading('EDUCATION',level=1)
    doc.add_paragraph('MS Management, Institute of Business Administration Karachi — ongoing since August 2025')
    doc.add_paragraph('BA (Honours) Communication & Design, minor in Philosophy, Habib University — June 2019')
    doc.add_heading('CERTIFICATIONS',level=1)
    doc.add_paragraph('Google Data Analytics Certificate | EU GSP+ Human Rights Compliance Reporting')
    doc.add_paragraph('APPLICATION DRAFT — Review for completeness and accuracy against your full master CV before sending.')
    return docx_bytes(doc)

def make_letter_docx(job,profile,contact=None):
    contact=contact or {};doc=Document();setup_doc(doc)
    doc.add_heading('Cover letter',0)
    doc.add_paragraph((contact.get('name') or 'Syed Mohammad Hassan')+' | '+(contact.get('email') or 'Contact details to add'))
    doc.add_paragraph('Re: '+job['title']+' — '+job['company'])
    doc.add_paragraph('Dear Hiring Manager,')
    doc.add_paragraph('I am applying for the '+job['title']+' role at '+job['company']+'. I offer experience in regional corporate communications, brand strategy and organisational change across international markets.')
    facts=relevant_facts(job,profile,limit=4)
    for f in facts[:3]:
        sentence=f[0].lower()+f[1:] if f and f[0].isupper() else f
        doc.add_paragraph('In my previous roles, I have '+sentence.rstrip('.').rstrip()+' .'.replace(' .','.'))
    doc.add_paragraph('My background combines hands-on communications delivery with stakeholder engagement and structured performance reporting. I would welcome the opportunity to discuss how this experience could contribute to your team.')
    doc.add_paragraph('I am currently based in Pakistan and would require employer-supported work authorisation to relocate. I would appreciate confirmation of the sponsorship options applicable to this vacancy.')
    doc.add_paragraph('Yours faithfully,\n'+(contact.get('name') or 'Syed Mohammad Hassan'))
    doc.add_paragraph('DRAFT — Verify facts, add contact details and edit the letter before approval.')
    return docx_bytes(doc)

def make_evidence_docx(job,evidence):
    doc=Document();setup_doc(doc);doc.add_heading('Sponsorship evidence',0)
    doc.add_paragraph(job['company']+' — '+job['title']);doc.add_paragraph(job.get('url',''))
    assessment,_=sponsorship_assessment(evidence);doc.add_paragraph('Assessment: '+assessment)
    for ev in evidence:
        doc.add_heading(ev['kind']+' | '+ev['verdict'],level=2)
        doc.add_paragraph(ev['quote']);doc.add_paragraph('Source: '+ev['source_url'])
        doc.add_paragraph('Recorded: '+ev['checked_at']+' | '+ev.get('reviewer',''))
    doc.add_paragraph('Evidence is a research aid, not a legal eligibility determination. Check current immigration rules and confirm with the employer.')
    return docx_bytes(doc)
